"""Tests for email normalization and duplicate grouping."""

from find_dupes import find_duplicates, normalize_email


class TestNormalizeEmail:
    """Canonicalization rules for individual addresses."""

    def test_lowercases_address(self):
        assert normalize_email("Foo@Gmail.COM") == "foo@gmail.com"

    def test_strips_surrounding_whitespace(self):
        assert normalize_email("  foo@example.com ") == "foo@example.com"

    def test_strips_dots_for_gmail(self):
        assert normalize_email("j.oh.n@gmail.com") == "john@gmail.com"

    def test_strips_dots_for_googlemail(self):
        assert normalize_email("j.ohn@googlemail.com") == "john@googlemail.com"

    def test_keeps_dots_for_other_domains(self):
        assert normalize_email("jo.hn@somecompany.com") == "jo.hn@somecompany.com"
        assert normalize_email("jo.hn@outlook.com") == "jo.hn@outlook.com"

    def test_strips_plus_alias(self):
        assert normalize_email("john+spam@gmail.com") == "john@gmail.com"
        assert normalize_email("john+a+b@example.com") == "john@example.com"

    def test_plus_alias_with_dots_on_gmail(self):
        assert normalize_email("j.ohn+news.letter@gmail.com") == "john@gmail.com"

    def test_invalid_emails_return_none(self):
        assert normalize_email("not-an-email") is None
        assert normalize_email("a@b@c.com") is None
        assert normalize_email("@example.com") is None
        assert normalize_email("user@") is None
        assert normalize_email("") is None


class TestFindDuplicates:
    """Grouping behavior across a whole subscriber list."""

    def test_groups_equivalent_addresses(self):
        emails = [
            "j.ohn@gmail.com",
            "john@gmail.com",
            "john+x@gmail.com",
            "unrelated@example.com",
        ]
        assert find_duplicates(emails) == [
            (
                "john@gmail.com",
                ["j.ohn@gmail.com", "john@gmail.com", "john+x@gmail.com"],
            ),
        ]

    def test_case_insensitive_match(self):
        emails = ["Foo@Example.com", "foo@example.com"]
        assert find_duplicates(emails) == [
            ("foo@example.com", ["Foo@Example.com", "foo@example.com"]),
        ]

    def test_no_duplicates_returns_empty(self):
        assert find_duplicates(["a@example.com", "b@example.com"]) == []

    def test_skips_unparseable_emails(self):
        emails = ["bogus", "john@gmail.com", "j.ohn@gmail.com"]
        assert find_duplicates(emails) == [
            ("john@gmail.com", ["john@gmail.com", "j.ohn@gmail.com"]),
        ]

    def test_dots_at_non_gmail_domains_are_distinct(self):
        assert find_duplicates(["jo.hn@somecompany.com", "john@somecompany.com"]) == []

    def test_groups_sorted_by_normalized_address(self):
        emails = [
            "z@example.com",
            "z+1@example.com",
            "a@example.com",
            "a+1@example.com",
        ]
        result = find_duplicates(emails)
        assert [normalized for normalized, _ in result] == [
            "a@example.com",
            "z@example.com",
        ]


class TestSaveExcel:
    """Spreadsheet output round-trip."""

    def test_writes_header_and_groups(self, tmp_path):
        from openpyxl import load_workbook

        from save_excel import save_excel

        out = tmp_path / "report.xlsx"
        save_excel([("john@gmail.com", ["j.ohn@gmail.com", "john+x@gmail.com"])], out)

        rows = list(load_workbook(out).active.iter_rows(values_only=True))
        assert rows == [
            ("Normalized Email", "Duplicate Subscribers"),
            ("john@gmail.com", "j.ohn@gmail.com, john+x@gmail.com"),
        ]
